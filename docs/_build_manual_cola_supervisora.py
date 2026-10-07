# -*- coding: utf-8 -*-
"""Genera el manual de cola para supervisión (Word)."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("manual-cola-mis-clientes-supervisora.docx")

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
TEAL = RGBColor(0x0F, 0x4C, 0x5C)
DARK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "1E3A5F"
ROW_ALT = "F1F5F9"


def set_run_font(run, *, size=11, bold=False, color=DARK, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CBD5E1")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, *, bold=False, size=10, color=DARK, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_cell_border(cell)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=NAVY if level == 1 else TEAL)
    h.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    return h


def p(doc, text, *, size=11, bold=False, space_after=8):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # Support **bold** segments
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = para.add_run(part)
        set_run_font(run, size=size, bold=bold or (i % 2 == 1), color=DARK)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.left_indent = Cm(1.0)
    # clear default run and add formatted
    if para.runs:
        para.runs[0].text = ""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = para.add_run(part)
        set_run_font(run, size=11, bold=(i % 2 == 1), color=DARK)
    return para


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell_text(cell, h, bold=True, size=10, color=WHITE)
        shade_cell(cell, HEADER_BG)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell_text(cell, val, size=10, bold=c_i == 0)
            if r_i % 2 == 1:
                shade_cell(cell, ROW_ALT)
        prevent_row_split(table.rows[r_i + 1])
    prevent_row_split(table.rows[0])
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("Manual de cola — Mis clientes  ·  Optick CRM  ·  Confidencial interno  ·  Pág. ")
    set_run_font(r1, size=8, color=MUTED)
    add_page_number(fp)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Capacitación de agentes  ·  octubre 2026")
    set_run_font(hr, size=8, color=MUTED)

    # ——— Portada ———
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kr = kicker.add_run("OPTICK CRM")
    set_run_font(kr, size=13, bold=True, color=TEAL)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("Manual de cola — Mis clientes")
    set_run_font(tr, size=28, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Guía para supervisión y capacitación de agentes")
    set_run_font(sr, size=14, color=MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("Octubre 2026")
    set_run_font(mr, size=12, bold=True, color=TEAL)

    box = doc.add_paragraph()
    box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    box.paragraph_format.space_before = Pt(24)
    br = box.add_run(
        "Este documento explica cómo trabaja el agente en Mis clientes:\n"
        "qué empresa llamar, cómo guardar, a dónde va cada respuesta\n"
        "y qué no debe hacer. Léalo con el sistema abierto."
    )
    set_run_font(br, size=11, color=DARK)

    doc.add_page_break()

    # ——— Índice numerado (visible al abrir / imprimir) ———
    add_heading_styled(doc, "Índice", 1)
    p(
        doc,
        "Use este índice para ubicar el tema. En Word también puede actualizar un campo de tabla de contenido (clic derecho → Actualizar campos) si lo agrega al documento.",
        size=10,
    )
    index_items = [
        ("1.", "Para qué sirve Mis clientes"),
        ("2.", "Entel vs Movistar"),
        ("3.", "Lote y FIFO (terminar el lote actual)"),
        ("4.", "Empresa vs contacto"),
        ("5.", "Guardar, siguiente empresa y siguiente pendiente"),
        ("6.", "Modal de segundo contacto"),
        ("7.", "Modal al cambiar respuesta del mismo contacto"),
        ("8.", "Respuestas y porcentajes"),
        ("9.", "Dónde va cada respuesta"),
        ("10.", "Lista, Detalle y Tarjetas"),
        ("11.", "Tras guardar: Resultado guardado vs Actualizando cola"),
        ("12.", "Qué no hacer"),
        ("13.", "Preguntas frecuentes"),
    ]
    for num, label in index_items:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.space_before = Pt(0)
        r_n = para.add_run(f"{num}  ")
        set_run_font(r_n, size=12, bold=True, color=TEAL)
        r_l = para.add_run(label)
        set_run_font(r_l, size=12, color=DARK)

    # TOC field (Word can refresh it)
    toc_p = doc.add_paragraph()
    toc_p.paragraph_format.space_before = Pt(12)
    run = toc_p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-2" \h \z \u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "(Campo de índice de Word: clic derecho → Actualizar campos)"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(hint)
    r.append(fld_end)

    doc.add_page_break()

    # 1
    add_heading_styled(doc, "1. Para qué sirve Mis clientes", 1)
    p(
        doc,
        "Mis clientes es la pantalla de trabajo del agente. Ahí ve las empresas que le asignaron y registra cada llamada. No es un listado de toda la empresa: es su cola personal.",
    )
    p(
        doc,
        "El objetivo es simple: llamar, anotar qué pasó y pasar a la siguiente empresa. El sistema guarda el resultado y ordena la cola para que usted no pierda el hilo.",
    )
    bullet(doc, "Arriba elige el **operador** (Entel o Movistar) y el **lote**.")
    bullet(doc, "El conmutador **Detalle** / **Tarjetas** / **Lista** cambia cómo ve las empresas, no cambia a quién le toca llamar.")
    bullet(doc, "En **Detalle** trabaja ficha por ficha (1 de N). Esa es la cola de trabajo.")

    # 2
    add_heading_styled(doc, "2. Entel vs Movistar", 1)
    p(
        doc,
        "El agente trabaja **un operador a la vez**. Entel y Movistar no se mezclan: cada archivo de importación, cada lote y cada cola son de un solo operador.",
    )
    bullet(doc, "Antes de llamar, confirme que el botón de operador coincide con la campaña del día.")
    bullet(doc, "Si tiene una cita del otro operador, el sistema puede avisarle. Cambie de operador solo cuando corresponda; no mezcle llamadas de ambos en la misma sesión de trabajo.")
    bullet(doc, "El lote FIFO (el orden de los archivos) es **solo de ese operador**.")

    # 3
    add_heading_styled(doc, "3. Lote y FIFO (terminar el lote actual)", 1)
    p(
        doc,
        "Un **lote** es un archivo de empresas que se importó y se le asignó. Por defecto la cola sigue **FIFO: más antiguo primero**. Eso significa: termine el lote que está en curso antes de saltar a uno más nuevo.",
    )
    bullet(doc, "El selector de lotes muestra cuántas empresas quedan pendientes (el termómetro de color).")
    bullet(doc, "Usted puede cambiar de lote a mano si supervisión se lo indica. No lo haga por costumbre.")
    bullet(doc, "En **Gestión de agentes**, supervisión puede poner **Más reciente primero** o **Todos los lotes**. El valor por defecto es terminar el actual.")
    p(
        doc,
        "Regla de oro: si el lote todavía tiene pendientes, siga ahí. Saltar de lote deja empresas a medias y desordena el trabajo del equipo.",
    )

    # 4
    add_heading_styled(doc, "4. Empresa vs contacto", 1)
    p(
        doc,
        "La **empresa** se identifica por el **RUC** (razón social). Cada empresa puede tener varios **contactos** (pestañas con nombre y teléfono).",
    )
    add_table(
        doc,
        ["Concepto", "Qué es", "Cómo se cambia"],
        [
            [
                "Empresa",
                "Un RUC. Es la ficha de la cola (1 de N).",
                "Botones siguiente / anterior, o clic en Lista.",
            ],
            [
                "Contacto",
                "Una persona o número de esa empresa.",
                "A mano: usted elige la pestaña. El sistema no salta de contacto solo.",
            ],
            [
                "Siguiente empresa",
                "Otro RUC, otra ficha.",
                "No es el siguiente teléfono de la misma empresa.",
            ],
        ],
        col_widths=[4.0, 6.5, 6.5],
    )
    p(
        doc,
        "Si la empresa tiene tres contactos, las pestañas quedan de izquierda a derecha **como en el Excel** (orden de la plantilla) y no se reordenan. Usted decide a cuál llama. Al **Guardar resultado**, la pestaña se pone **verde en su mismo lugar** (1 de 3 sigue 1 de 3). **Guardar resultado** / **Guardar actualización** no cambia de persona: se queda en la pestaña actual. El cambio de contacto es solo a mano (pestañas o flechas). Cada pestaña muestra su última **Respuesta**. Cambia de **empresa** solo si usted pulsa **Guardar y siguiente empresa** o **Guardar y siguiente pendiente**.",
    )

    # 5
    add_heading_styled(doc, "5. Guardar, siguiente empresa y siguiente pendiente", 1)
    p(doc, "Hay tres botones verdes/azules en la parte baja de **Detalle**:")
    add_table(
        doc,
        ["Botón", "Qué hace", "Cuándo usarlo"],
        [
            [
                "Guardar resultado",
                "Graba la respuesta en el contacto actual y se queda en la misma persona (misma pestaña, mismo 1 de 3). La pestaña se pone verde en su sitio; no pasa al final. Si ya había un registro, el botón dice Guardar actualización. Cada pestaña muestra su última respuesta.",
                "Cuando quiere revisar o corregir sin saltar de persona ni de empresa.",
            ],
            [
                "Guardar y siguiente empresa",
                "Graba y pasa al siguiente RUC de la cola (el que sigue en el orden, tenga o no registro).",
                "Cuando ya terminó con esa empresa y quiere la siguiente de la fila.",
            ],
            [
                "Guardar y siguiente pendiente",
                "Graba y salta a la próxima empresa que todavía no tiene registro suyo en este lote.",
                "Cuando quiere ir directo a las que aún no ha tocado.",
            ],
        ],
        col_widths=[5.0, 6.5, 5.5],
    )
    p(
        doc,
        "El botón verde **Guardar resultado** / **Guardar actualización** no cambia de persona: las pestañas (o las flechas) son el único cambio de contacto. Al abrir otra pestaña, **Respuesta** y las notas son las de ese contacto (o **— Seleccionar —** si aún no tiene registro suyo). **Guardar y siguiente empresa** y **Guardar y siguiente pendiente** sí cambian de **RUC**.",
    )
    p(
        doc,
        "Debe elegir una **Respuesta** antes de guardar. Si no hay teléfono de usuario, el sistema no deja grabar el resultado de la llamada. Complete el teléfono o pase al contacto que sí lo tiene.",
    )

    # 6
    add_heading_styled(doc, "6. Modal de segundo contacto", 1)
    p(
        doc,
        "Si la empresa **ya tiene una llamada guardada en otro contacto** y usted intenta guardar en uno distinto, aparece un aviso. Es válido: por ejemplo **No contesta** en más de un número de la misma empresa.",
    )
    bullet(doc, "Lea el nombre del contacto que ya registró y el que va a registrar ahora.")
    bullet(doc, "Confirme si corresponde. Puede marcar **no volver a preguntar** en esa empresa (solo durante la sesión).")
    bullet(doc, "Si se equivocó de pestaña, cancele y cambie de contacto a mano.")

    # 7
    add_heading_styled(doc, "7. Modal al cambiar respuesta del mismo contacto", 1)
    p(
        doc,
        "Si **este contacto** ya tiene una respuesta suya y usted elige **otra respuesta** distinta, aparece un aviso **antes** de Guardando…. No es el modal de segundo contacto: es el mismo número, otra etiqueta.",
    )
    bullet(doc, "Título: **Este contacto ya tiene una respuesta**.")
    bullet(
        doc,
        "Le dice qué registró antes y qué va a guardar ahora. La anterior **queda en el historial**; no se borra.",
    )
    bullet(doc, "Puede **Cancelar** o pulsar **Sí, guardar nueva respuesta**.")
    bullet(
        doc,
        "Opcional: **No volver a preguntar en este contacto** (solo esta sesión, por ese contacto).",
    )
    p(
        doc,
        "No aparece en el primer guardado, ni si solo cambia notas, ni si solo mueve la agenda, ni si la respuesta es la misma.",
    )

    # 8
    add_heading_styled(doc, "8. Respuestas y porcentajes", 1)
    p(
        doc,
        "Cada respuesta tiene un porcentaje. Eso no es una nota del agente: es el avance de la empresa en el embudo.",
    )
    add_table(
        doc,
        ["Grupo", "Respuestas", "Porcentaje"],
        [
            [
                "Operativas (0%)",
                "No contesta, Volver a llamar, Sin llegada al decisor, RUC suspendido, Cliente actual, No interesado",
                "0%",
            ],
            [
                "Embudo comercial",
                "Interesado → Propuesta presentada → Discusión de propuesta → A la espera de respuesta final",
                "25% · 50% · 75% · 90%",
            ],
            ["Cierre de venta", "Venta cerrada", "100%"],
        ],
        col_widths=[4.5, 9.5, 3.0],
    )
    p(
        doc,
        "**Volver a llamar** exige fecha y hora de agenda. **No interesado** y **Venta cerrada** no se agendan. El resto puede llevar agenda si supervisión lo pide.",
    )

    # 9
    add_heading_styled(doc, "9. Dónde va cada respuesta", 1)
    p(
        doc,
        "Esto es lo que el agente debe memorizar. **Detalle** es la cola de trabajo (1 de N). **Lista** puede mostrar más con filtros. **Otros** es el archivo de respuestas 0% que no son No contesta.",
    )
    add_table(
        doc,
        ["Respuesta", "¿Sigue en Detalle?", "¿Dónde más?", "Nota para el agente"],
        [
            [
                "Pendiente (sin respuesta)",
                "Sí",
                "Lista · filtro Pendiente",
                "Aún no hay registro. Es el trabajo nuevo.",
            ],
            [
                "Volver a llamar + agenda",
                "Sí",
                "Lista · Volver a llamar · panel Agendados",
                "Debe cargar fecha y hora. Cumpla la cita.",
            ],
            [
                "No contesta (1 en la empresa)",
                "Sí",
                "Lista · cola No contesta",
                "Cuenta solo No contesta, no cualquier llamada. Un No contesta no saca la ficha.",
            ],
            [
                "No contesta (2 o más en la empresa)",
                "No",
                "Lista · No contesta — depurado",
                "Dos No contesta del agente en esa empresa. Sale de Detalle. Nunca va a Otros.",
            ],
            [
                "Sin llegada al decisor",
                "Sí",
                "Lista · Otros",
                "Es un registro guardado: el agente puede volver a llamar. Está en Detalle y también en Otros.",
            ],
            [
                "No interesado",
                "No",
                "Lista · Otros",
                "Cierre. Fuera de Detalle. Solo archivo en Otros.",
            ],
            [
                "Cliente actual",
                "No",
                "Lista · Otros",
                "Cierre. Fuera de Detalle. Solo archivo en Otros.",
            ],
            [
                "RUC suspendido / no habido",
                "No",
                "Lista · Otros",
                "Cierre. Fuera de Detalle. Solo archivo en Otros.",
            ],
            [
                "Interesado … Venta cerrada",
                "Sí",
                "Lista · chips del embudo",
                "Avance comercial 25% a 100%. Siguen en la cola de trabajo.",
            ],
        ],
        col_widths=[4.2, 3.2, 4.4, 5.2],
    )
    p(
        doc,
        "Resumen en una frase: **Sin llegada** se queda en **Detalle**. **No interesado**, **cliente actual** y **RUC suspendido** salen de **Detalle** y viven en **Otros**. **Dos No contesta** en la **empresa** (no dos llamadas cualesquiera) van a **depurado**, no a Otros.",
    )
    p(
        doc,
        "Ejemplo: **Sin llegada** y después **No contesta** (Guardar actualización crea una fila nueva) = **un** No contesta. Sigue en **Detalle** y en la cola **No contesta**. No entra a depurado.",
    )

    # 10
    add_heading_styled(doc, "10. Lista, Detalle y Tarjetas", 1)
    p(doc, "Tres vistas, el mismo trabajo:")
    bullet(doc, "**Detalle:** ficha completa, historial, agenda y botones de guardar. Es donde se trabaja.")
    bullet(doc, "**Tarjetas:** grilla con búsqueda. Útil para ubicar rápido.")
    bullet(doc, "**Lista:** tabla de empresas. Filtros de cola (Todos, embudo, Pendiente, Volver a llamar, No contesta, No contesta — depurado, Otros).")
    add_heading_styled(doc, "Filtro Otros", 2)
    p(
        doc,
        "**Otros** lista el resto 0% **excepto No contesta**: Sin llegada, No interesado, cliente actual, RUC suspendido y similares. Sin llegada aparece en **Detalle** y también en **Otros**. No contesta nunca entra en Otros (tiene sus propias colas).",
    )
    add_heading_styled(doc, "Por qué a veces el # es un guion (—)", 2)
    p(
        doc,
        "La columna **#** es el número en la cola de **Detalle** (1, 2, 3…). Si la empresa **no está** en esa cola —por ejemplo No interesado, cliente actual, RUC suspendido o No contesta depurado— el # muestra **—**. Eso no significa que la fila esté vacía: significa “fuera de la cola de trabajo”.",
    )
    p(
        doc,
        "Al hacer clic en esa fila, el sistema abre **esa misma empresa** (ese RUC). El encabezado puede decir **Fuera de la cola**. Use **Volver a la lista** para regresar. No debe aparecer otra empresa (por ejemplo Agrotours) en lugar de la que usted eligió.",
    )

    # 11
    add_heading_styled(doc, "11. Tras guardar: Resultado guardado vs Actualizando cola", 1)
    p(
        doc,
        "Cuando pulsa **Guardar**, el resultado se graba enseguida. Verá **Resultado guardado** (o el aviso de No contesta — depurado si era el segundo **No contesta** de esa empresa). El botón se libera: usted **puede seguir trabajando**.",
    )
    p(
        doc,
        "Al mismo tiempo puede aparecer **Actualizando cola…**. Eso es el sistema reordenando la lista en segundo plano. No es un error y no tiene que esperar a que desaparezca para seguir. No machaque el botón Guardar.",
    )

    # 12
    add_heading_styled(doc, "12. Qué no hacer", 1)
    bullet(doc, "**No martille Guardar.** Un clic basta. Si el botón dice Guardando…, espere el Resultado guardado.")
    bullet(doc, "**No resetee la campaña** ni pida borrar lotes para “empezar de nuevo” sin autorización de supervisión.")
    bullet(doc, "**No mezcle operadores.** No trabaje Entel y Movistar a la vez ni importe un archivo con los dos nombres.")
    bullet(doc, "**No confunda siguiente empresa con siguiente contacto.** Siguiente empresa = otro RUC. El contacto se cambia a mano en las pestañas.")
    bullet(doc, "**No busque No interesado ni RUC suspendido en Detalle.** Están en **Lista → Otros**.")
    bullet(doc, "**No busque No contesta (2 veces en la empresa) en Otros.** Está en **No contesta — depurado**.")
    bullet(doc, "**No salte de lote** si el actual todavía tiene pendientes, salvo indicación de supervisión.")

    # 13
    add_heading_styled(doc, "13. Preguntas frecuentes", 1)

    add_heading_styled(doc, "Hice clic en Lista y se abrió otra empresa (Agrotours u otra). ¿Es normal?", 2)
    p(
        doc,
        "No. Ese era un error cuando el # era **—**: el sistema saltaba a la ficha que ya tenía abierta (otro RUC). Ya está corregido: el clic abre **la misma empresa** que usted eligió. Si el # es —, verá **Fuera de la cola** y **Volver a la lista**.",
    )

    add_heading_styled(doc, "Guardé Sin llegada y la empresa desapareció de Detalle. ¿Volverá?", 2)
    p(
        doc,
        "No debe desaparecer. **Sin llegada** es un registro guardado y **sigue en Detalle** para que el agente pueda volver a llamar. También la encuentra en **Lista → Otros**.",
    )

    add_heading_styled(doc, "¿Por qué No contesta a veces está en la cola y a veces no?", 2)
    p(
        doc,
        "Depurado cuenta **solo No contesta** de esa **empresa** (sus registros, no dos llamadas cualesquiera). Un No contesta: sigue en **Detalle** y en la cola **No contesta**. Dos o más No contesta: sale de Detalle y va a **No contesta — depurado**. Nunca a Otros.",
    )

    add_heading_styled(doc, "Guardé Sin llegada y luego No contesta. ¿Se fue a depurado?", 2)
    p(
        doc,
        "No. Eso es **un** No contesta. Aunque Guardar actualización cree una fila nueva, la anterior no era No contesta. La ficha sigue en **Detalle** hasta que haya **dos No contesta** en la empresa.",
    )

    add_heading_styled(doc, "Cambié la respuesta del mismo contacto y salió un aviso. ¿Borré la anterior?", 2)
    p(
        doc,
        "No. La respuesta anterior queda en el historial. El aviso solo confirma que va a guardar una **nueva**. Si solo cambió notas o la agenda, ese aviso no aparece.",
    )

    add_heading_styled(doc, "¿Puedo llamar de nuevo a un No interesado?", 2)
    p(
        doc,
        "La ficha ya no está en la cola de trabajo. Supervisión puede abrirla desde **Lista → Otros**. El agente no debe “buscarla” avanzando 1/N: no está ahí.",
    )

    add_heading_styled(doc, "¿El porcentaje lo elijo yo?", 2)
    p(
        doc,
        "No. Usted elige la **Respuesta**. El porcentaje lo pone el sistema (0% operativas, 25% a 100% el embudo).",
    )

    add_heading_styled(doc, "¿Qué hago si hay dos contactos y ya llamé a uno?", 2)
    p(
        doc,
        "Cambie de pestaña a mano. El orden de las pestañas es el del Excel y no cambia. Cada pestaña muestra su última respuesta; **Guardar resultado** no salta a otra persona ni mueve el verde al final. Si guarda en el segundo, el aviso de segundo contacto le pedirá confirmar. Es correcto registrar más de un número de la misma empresa.",
    )

    add_heading_styled(doc, "Guardé y me cambió de persona (otra pestaña). ¿Es normal?", 2)
    p(
        doc,
        "No. El botón verde se queda en el mismo contacto y en el **mismo lugar** (1 de 3 sigue 1 de 3). Las pestañas no se reordenan: el verde no pasa al final. Si ve la **Respuesta** de Mario en la pestaña de otra persona, pulse esa pestaña de nuevo: debe cargar la respuesta de ese contacto (o **— Seleccionar —** si no tiene registro). El cambio de persona es solo con las pestañas o las flechas.",
    )

    add_heading_styled(doc, "Guardé y la pestaña verde se fue al final. ¿Es normal?", 2)
    p(
        doc,
        "No. Las pestañas quedan de izquierda a derecha como en el Excel. Al guardar, esa pestaña se pone verde **en su sitio**. El orden de la plantilla no cambia.",
    )

    add_heading_styled(doc, "Cierre para quien capacita", 2)
    p(
        doc,
        "Pida al agente que, con el sistema abierto, marque una empresa pendiente, una Sin llegada, un No interesado (Otros) y un No contesta de dos No contesta en la empresa (depurado). Si distingue esas cuatro rutas, ya entiende la cola.",
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    nr = note.add_run(
        "Documento interno · Optick CRM · Manual de cola — Mis clientes · octubre 2026. "
        "Pensado para imprimir y para capacitación en sala. No sustituye las Novedades del sistema."
    )
    set_run_font(nr, size=9, color=MUTED)

    try:
        doc.save(OUT)
    except PermissionError:
        tmp = OUT.with_name(OUT.stem + ".tmp.docx")
        doc.save(tmp)
        print(f"LOCKED {OUT.resolve()}")
        print(tmp.resolve())
        raise SystemExit(1)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
